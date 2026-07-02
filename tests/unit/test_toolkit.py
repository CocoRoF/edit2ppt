"""Tests for the agent-toolkit surfaces: facade, agent_tools, local MCP.

LLM-backed verbs (generate/edit) are exercised with a stubbed client via
the existing edit_deck stubs elsewhere; here we cover the deterministic
verbs end-to-end plus schema/dispatch/lazy-import contracts.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

import edit2ppt
from edit2ppt import analyze_pptx, preview_pptx, run_tool, set_pptx_text
from edit2ppt.agent_tools import ANTHROPIC_TOOLS, TOOL_NAMES


@pytest.fixture
def deck_path(tmp_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "원래 제목"
    gf = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(2))
    gf.table.cell(0, 0).text = "셀 텍스트"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


class TestLazyPackageSurface:
    def test_version_and_lazy_exports(self):
        assert edit2ppt.__version__
        assert callable(edit2ppt.generate_pptx)
        assert callable(edit2ppt.edit_pptx)
        assert "ANTHROPIC_TOOLS" in dir(edit2ppt)

    def test_unknown_attribute_raises(self):
        with pytest.raises(AttributeError):
            edit2ppt.does_not_exist  # noqa: B018


class TestFacadeDeterministicVerbs:
    def test_preview_strings_and_files(self, deck_path, tmp_path):
        svgs = preview_pptx(deck_path)
        assert len(svgs) == 1 and "원래 제목" in svgs[0]
        paths = preview_pptx(deck_path, out_dir=tmp_path / "svg")
        assert [p.name for p in paths] == ["slide_000.svg"]
        assert "원래 제목" in Path(paths[0]).read_text(encoding="utf-8")

    def test_analyze_lists_addressable_paragraphs(self, deck_path):
        info = analyze_pptx(deck_path)
        assert info["page_count"] == 1
        texts = info["slides"][0]["texts"]
        shape_entries = [t for t in texts if "shape_id" in t]
        table_entries = [t for t in texts if "table_id" in t]
        assert any(t["text"] == "원래 제목" for t in shape_entries)
        cell = next(t for t in table_entries if t["text"] == "셀 텍스트")
        assert (cell["row"], cell["col"]) == (0, 0)

    def test_set_text_via_analyze_addresses(self, deck_path, tmp_path):
        info = analyze_pptx(deck_path)
        target = next(
            t for t in info["slides"][0]["texts"] if t.get("text") == "원래 제목"
        )
        out = tmp_path / "out.pptx"
        result = set_pptx_text(
            deck_path,
            [
                {
                    "slide": 0,
                    "shape_id": target["shape_id"],
                    "para": target["para"],
                    "new_text": "파사드 수정",
                    "old_text": target["text"],
                }
            ],
            output=out,
        )
        assert result.applied == 1 and result.path == out
        assert "파사드 수정" in preview_pptx(out)[0]

    def test_zero_applied_leaves_input_untouched(self, deck_path, tmp_path):
        result = set_pptx_text(
            deck_path,
            [{"slide": 0, "shape_id": 99999, "para": 0, "new_text": "x"}],
            output=tmp_path / "never.pptx",
        )
        assert result.applied == 0
        assert not (tmp_path / "never.pptx").exists()
        assert result.path == deck_path


class TestAgentTools:
    def test_schemas_are_anthropic_shaped(self):
        assert TOOL_NAMES == [
            "generate_pptx", "edit_pptx", "preview_pptx", "set_pptx_text", "analyze_pptx",
        ]
        for tool in ANTHROPIC_TOOLS:
            assert set(tool) == {"name", "description", "input_schema"}
            schema = tool["input_schema"]
            assert schema["type"] == "object"
            assert schema["required"], tool["name"]

    def test_dispatch_deterministic_tools(self, deck_path, tmp_path):
        info = run_tool("analyze_pptx", {"pptx": str(deck_path)})
        assert info["page_count"] == 1
        res = run_tool(
            "preview_pptx", {"pptx": str(deck_path), "out_dir": str(tmp_path / "s")}
        )
        assert res["page_count"] == 1 and Path(res["svg_paths"][0]).exists()

    def test_unknown_tool_raises(self):
        with pytest.raises(ValueError, match="unknown edit2ppt tool"):
            run_tool("rm_rf_slash", {})


class TestLocalMcpServer:
    def test_tool_registry_matches_agent_tools(self):
        from edit2ppt.mcp.local_server import build_local_mcp_server

        mcp = build_local_mcp_server()
        tools = asyncio.run(mcp.list_tools())
        assert sorted(t.name for t in tools) == sorted(TOOL_NAMES)

    def test_deterministic_tool_call_through_mcp(self, deck_path):
        from edit2ppt.mcp.local_server import build_local_mcp_server

        mcp = build_local_mcp_server()
        result = asyncio.run(
            mcp.call_tool("analyze_pptx", {"pptx": str(deck_path)})
        )
        # FastMCP returns (content_blocks, raw_result)
        raw = result[1] if isinstance(result, tuple) else result
        assert "1" in str(raw) or "page_count" in str(raw)
